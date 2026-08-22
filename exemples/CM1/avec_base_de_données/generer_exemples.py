"""Genere plusieurs releves de compte factices, un par mois.

Ces fichiers servent a alimenter l'application : on les importe un par un
pour voir la liste se remplir et la comparaison des trois derniers releves
devenir interessante.

Usage :
    python generer_exemples.py [dossier_de_sortie]
"""

import os
import random
import sys
from datetime import date

from openpyxl import Workbook

# Les quatre mois generes : le dernier importe sera donc le plus recent.
MOIS_A_GENERER = [(2026, 2), (2026, 3), (2026, 4), (2026, 5)]

# Operations presentes tous les mois : (jour, libelle, categorie, montant)
OPERATIONS_MENSUELLES = [
    (2, "Loyer appartement", "Logement", -780.00),
    (3, "Assurance habitation", "Logement", -18.90),
    (5, "Abonnement internet", "Abonnements", -29.99),
    (5, "Forfait mobile", "Abonnements", -14.99),
    (6, "Abonnement streaming", "Abonnements", -11.99),
    (10, "Electricite - mensualite", "Logement", -62.00),
    (15, "Abonnement transports", "Transport", -75.20),
    (25, "Salaire", "Salaire", 2450.00),
]

# Depenses tirees au hasard : (libelle, categorie, montant min, montant max, nombre)
DEPENSES_ALEATOIRES = [
    ("Supermarche", "Courses", 18.0, 95.0, 8),
    ("Boulangerie", "Courses", 2.5, 12.0, 5),
    ("Restaurant", "Loisirs", 14.0, 48.0, 4),
    ("Cafe", "Loisirs", 2.0, 7.5, 4),
    ("Station essence", "Transport", 35.0, 72.0, 2),
    ("Pharmacie", "Sante", 6.0, 34.0, 2),
    ("Librairie", "Loisirs", 9.0, 27.0, 1),
]

# Evenements propres a un mois, pour que les releves ne se ressemblent pas.
EVENEMENTS = {
    (2026, 2): [(14, "Reparation voiture", "Transport", -410.00)],
    (2026, 3): [
        (8, "Prime annuelle", "Salaire", 850.00),
        (20, "Vacances - hotel", "Loisirs", -520.00),
    ],
    (2026, 4): [
        (12, "Lave-linge", "Maison", -639.00),
        (18, "Remboursement mutuelle", "Sante", 43.60),
    ],
    (2026, 5): [
        (22, "Vente vetements en ligne", "Revenus divers", 55.00),
        (27, "Consultation dentiste", "Sante", -120.00),
        (28, "Billet de train", "Transport", -89.40),
    ],
}


def construire_operations(annee, mois):
    """Retourne les operations d'un mois, triees par jour."""
    # Une graine differente par mois : les fichiers different, mais chaque
    # execution regenere exactement les memes.
    random.seed(annee * 100 + mois)

    operations = list(OPERATIONS_MENSUELLES) + EVENEMENTS.get((annee, mois), [])

    for libelle, categorie, mini, maxi, nombre in DEPENSES_ALEATOIRES:
        for _ in range(nombre):
            # Jusqu'a 28 seulement : ainsi la date existe dans tous les mois.
            jour = random.randint(1, 28)
            montant = -round(random.uniform(mini, maxi), 2)
            operations.append((jour, libelle, categorie, montant))

    operations.sort(key=lambda operation: operation[0])
    return operations


def ecrire_fichier(annee, mois, operations, chemin):
    """Ecrit les operations d'un mois dans un fichier Excel."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Releve"
    feuille.append(["Date", "Libelle", "Categorie", "Montant"])

    for jour, libelle, categorie, montant in operations:
        feuille.append([date(annee, mois, jour), libelle, categorie, montant])

    # "mm-dd-yy" est le format de date integre n0 14 du standard Excel : chaque
    # tableur l'affiche dans le format court de sa propre langue. Un code
    # personnalise comme "DD/MM/YYYY" serait affiche tel quel par certains
    # logiciels, dont Numbers sur Mac.
    for ligne in feuille.iter_rows(min_row=2, min_col=1, max_col=1):
        ligne[0].number_format = "mm-dd-yy"
    for ligne in feuille.iter_rows(min_row=2, min_col=4, max_col=4):
        ligne[0].number_format = "0.00"
    for colonne, largeur in zip("ABCD", (12, 28, 16, 12)):
        feuille.column_dimensions[colonne].width = largeur

    classeur.save(chemin)


def main():
    dossier = sys.argv[1] if len(sys.argv) > 1 else "donnees_exemple"
    os.makedirs(dossier, exist_ok=True)

    for annee, mois in MOIS_A_GENERER:
        operations = construire_operations(annee, mois)
        chemin = os.path.join(dossier, f"releve_{annee}_{mois:02d}.xlsx")
        ecrire_fichier(annee, mois, operations, chemin)
        print(f"{len(operations):3d} operations -> {chemin}")


if __name__ == "__main__":
    main()
