"""Lecture d'un releve Excel et enregistrement en base.

Le fichier attendu contient une ligne d'en-tete puis une ligne par operation :
Date, Libelle, Categorie, Montant. Les depenses sont negatives.
"""

import hashlib
import os

from openpyxl import load_workbook

from .base import convertir_en_date
from .modeles import Operation


class FichierInvalide(Exception):
    """Levee quand le fichier ne ressemble pas a un releve de compte."""


def lire_fichier_excel(chemin):
    """Retourne la liste des operations contenues dans le fichier.

    Les lignes vides ou dont le montant n'est pas un nombre sont ignorees :
    un vrai relevé bancaire contient souvent des lignes de titre ou de total.
    """
    try:
        classeur = load_workbook(chemin, data_only=True, read_only=True)
    except Exception as erreur:
        raise FichierInvalide(f"Fichier Excel illisible : {erreur}") from erreur

    feuille = classeur.active
    operations = []

    for ligne in feuille.iter_rows(min_row=2, max_col=4, values_only=True):
        valeur_date, libelle, categorie, montant = ligne
        if valeur_date is None or not isinstance(montant, (int, float)):
            continue
        try:
            jour = convertir_en_date(valeur_date)
        except (TypeError, ValueError):
            continue  # la colonne Date ne contient pas une date : on ignore
        operations.append(
            Operation(
                date=jour,
                libelle=str(libelle or "(sans libelle)").strip(),
                categorie=str(categorie or "Sans categorie").strip(),
                montant=float(montant),
            )
        )

    classeur.close()

    if not operations:
        raise FichierInvalide(
            "Aucune operation trouvee. Colonnes attendues : "
            "Date, Libelle, Categorie, Montant."
        )

    return operations


def empreinte_fichier(chemin):
    """Calcule l'empreinte SHA-256 du fichier.

    Deux fichiers au contenu identique ont la meme empreinte, meme s'ils ont
    ete renommes : c'est ce qui permet de detecter un import en double.
    """
    condensat = hashlib.sha256()
    with open(chemin, "rb") as fichier:
        for bloc in iter(lambda: fichier.read(65536), b""):
            condensat.update(bloc)
    return condensat.hexdigest()


def importer_releve(base, chemin):
    """Lit un fichier Excel et l'enregistre dans la base.

    Retourne (identifiant du releve, nombre d'operations importees).
    Leve `FichierInvalide` ou `ReleveDejaImporte` en cas de probleme.
    """
    operations = lire_fichier_excel(chemin)
    releve_id = base.ajouter_releve(
        nom_fichier=os.path.basename(chemin),
        empreinte=empreinte_fichier(chemin),
        operations=operations,
    )
    return releve_id, len(operations)
