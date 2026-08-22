"""Génère un relevé de compte factice au format Excel.

Le fichier produit sert de jeu de test pour `analyse_releve.py`.

Usage :
    python generer_exemple.py [fichier_de_sortie.xlsx]
"""

import random
import sys
from datetime import date

from openpyxl import Workbook

# Reproductibilité : le même fichier est généré à chaque exécution.
random.seed(42)

ANNEE = 2026
MOIS = 5  # mai

# Opérations fixes du mois : (jour, libellé, catégorie, montant)
OPERATIONS_FIXES = [
    (2, "Loyer appartement", "Logement", -780.00),
    (3, "Assurance habitation", "Logement", -18.90),
    (5, "Abonnement internet", "Abonnements", -29.99),
    (5, "Forfait mobile", "Abonnements", -14.99),
    (6, "Abonnement streaming", "Abonnements", -11.99),
    (10, "Electricite - mensualite", "Logement", -62.00),
    (15, "Abonnement transports", "Transport", -75.20),
    (18, "Remboursement mutuelle", "Sante", 43.60),
    (22, "Vente vetements en ligne", "Revenus divers", 55.00),
    (25, "Salaire mai", "Salaire", 2450.00),
    (27, "Consultation dentiste", "Sante", -120.00),
    (28, "Billet de train", "Transport", -89.40),
]

# Dépenses aléatoires : (libellé, catégorie, montant min, montant max, nombre)
DEPENSES_ALEATOIRES = [
    ("Supermarche", "Courses", 18.0, 95.0, 8),
    ("Boulangerie", "Courses", 2.5, 12.0, 5),
    ("Restaurant", "Loisirs", 14.0, 48.0, 4),
    ("Cafe", "Loisirs", 2.0, 7.5, 4),
    ("Station essence", "Transport", 35.0, 72.0, 2),
    ("Pharmacie", "Sante", 6.0, 34.0, 2),
    ("Librairie", "Loisirs", 9.0, 27.0, 1),
    ("Magasin de bricolage", "Maison", 12.0, 140.0, 1),
]


def construire_operations():
    """Retourne la liste des opérations du mois, triées par date."""
    operations = list(OPERATIONS_FIXES)

    for libelle, categorie, mini, maxi, nombre in DEPENSES_ALEATOIRES:
        for _ in range(nombre):
            jour = random.randint(1, 28)
            montant = -round(random.uniform(mini, maxi), 2)
            operations.append((jour, libelle, categorie, montant))

    operations.sort(key=lambda operation: operation[0])
    return operations


def ecrire_fichier(operations, chemin):
    """Écrit les opérations dans un fichier Excel."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Releve"

    feuille.append(["Date", "Libelle", "Categorie", "Montant"])

    for jour, libelle, categorie, montant in operations:
        feuille.append([date(ANNEE, MOIS, jour), libelle, categorie, montant])

    # Un peu de mise en forme pour que le fichier soit lisible.
    # "mm-dd-yy" est le format de date integre n0 14 du standard Excel : chaque
    # tableur l'affiche dans le format court de sa propre langue (02/05/2026 en
    # francais). Un code personnalise comme "DD/MM/YYYY" serait affiche tel quel
    # par certains logiciels, dont Numbers sur Mac.
    for ligne in feuille.iter_rows(min_row=2, min_col=1, max_col=1):
        ligne[0].number_format = "mm-dd-yy"
    for ligne in feuille.iter_rows(min_row=2, min_col=4, max_col=4):
        ligne[0].number_format = "0.00"
    for colonne, largeur in zip("ABCD", (12, 28, 16, 12)):
        feuille.column_dimensions[colonne].width = largeur

    classeur.save(chemin)


def main():
    chemin = sys.argv[1] if len(sys.argv) > 1 else "releve_mai_2026.xlsx"
    operations = construire_operations()
    ecrire_fichier(operations, chemin)
    print(f"{len(operations)} operations ecrites dans {chemin}")


if __name__ == "__main__":
    main()
