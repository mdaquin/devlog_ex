"""Analyse d'un relevé de compte mensuel au format Excel.

L'utilisateur choisit un fichier, le programme affiche :
  - le total des revenus,
  - le total des dépenses,
  - le solde du mois,
  - les 3 plus grosses dépenses.

Le fichier Excel doit contenir une ligne d'en-tête puis une ligne par
opération, avec les colonnes : Date, Libelle, Categorie, Montant.
Les dépenses sont des montants négatifs, les revenus des montants positifs.

Usage :
    python analyse_releve.py
"""

import tkinter as tk
from dataclasses import dataclass
from datetime import date, datetime
from tkinter import filedialog, messagebox

from openpyxl import load_workbook


@dataclass
class Operation:
    """Une ligne du relevé de compte."""

    date: date
    libelle: str
    categorie: str
    montant: float


def lire_releve(chemin):
    """Lit un fichier Excel et retourne la liste des opérations.

    Les lignes vides ou dont le montant n'est pas un nombre sont ignorées.
    """
    classeur = load_workbook(chemin, data_only=True)
    feuille = classeur.active

    operations = []
    # min_row=2 : on saute la ligne d'en-tête.
    for ligne in feuille.iter_rows(min_row=2, max_col=4, values_only=True):
        valeur_date, libelle, categorie, montant = ligne
        if not isinstance(montant, (int, float)):
            continue
        if isinstance(valeur_date, datetime):
            valeur_date = valeur_date.date()
        operations.append(
            Operation(
                date=valeur_date,
                libelle=str(libelle or "(sans libelle)"),
                categorie=str(categorie or ""),
                montant=float(montant),
            )
        )

    if not operations:
        raise ValueError("Aucune operation trouvee dans ce fichier.")

    return operations


def calculer_statistiques(operations):
    """Calcule les statistiques du mois à partir des opérations."""
    revenus = [op for op in operations if op.montant > 0]
    depenses = [op for op in operations if op.montant < 0]

    total_revenus = sum(op.montant for op in revenus)
    total_depenses = sum(op.montant for op in depenses)

    # La plus grosse dépense est le montant le plus négatif.
    plus_grosses_depenses = sorted(depenses, key=lambda op: op.montant)[:3]

    return {
        "nombre_operations": len(operations),
        "total_revenus": total_revenus,
        "total_depenses": total_depenses,
        "solde": total_revenus + total_depenses,
        "plus_grosses_depenses": plus_grosses_depenses,
    }


def formater_montant(montant):
    """Formate un montant en euros, par exemple : 1 234,56 EUR."""
    texte = f"{abs(montant):,.2f}"
    texte = texte.replace(",", " ").replace(".", ",")
    signe = "-" if montant < 0 else ""
    return f"{signe}{texte} EUR"


def formater_statistiques(stats):
    """Transforme les statistiques en texte affichable."""
    lignes = [
        f"Operations analysees : {stats['nombre_operations']}",
        "",
        f"Total des revenus  : {formater_montant(stats['total_revenus'])}",
        f"Total des depenses : {formater_montant(stats['total_depenses'])}",
        f"Solde du mois      : {formater_montant(stats['solde'])}",
        "",
        "3 plus grosses depenses :",
    ]

    for numero, operation in enumerate(stats["plus_grosses_depenses"], start=1):
        jour = operation.date.strftime("%d/%m/%Y") if operation.date else "??/??/????"
        lignes.append(
            f"  {numero}. {jour}  {operation.libelle} "
            f"({operation.categorie}) : {formater_montant(operation.montant)}"
        )

    if not stats["plus_grosses_depenses"]:
        lignes.append("  (aucune depense ce mois-ci)")

    return "\n".join(lignes)


class Application(tk.Tk):
    """Fenêtre principale : un bouton pour choisir le fichier, puis le résultat."""

    def __init__(self):
        super().__init__()
        self.title("Analyse d'un releve de compte")
        self.geometry("560x340")

        tk.Button(
            self,
            text="Choisir un releve Excel...",
            command=self.ouvrir_fichier,
        ).pack(pady=15)

        self.zone_resultat = tk.Label(
            self,
            text="Aucun fichier charge pour le moment.",
            justify="left",
            anchor="nw",
            font=("Courier", 12),
        )
        self.zone_resultat.pack(fill="both", expand=True, padx=20, pady=10)

    def ouvrir_fichier(self):
        chemin = filedialog.askopenfilename(
            title="Choisir un releve de compte",
            filetypes=[("Fichiers Excel", "*.xlsx"), ("Tous les fichiers", "*.*")],
        )
        if not chemin:  # l'utilisateur a annulé
            return

        try:
            operations = lire_releve(chemin)
            stats = calculer_statistiques(operations)
            resultat = formater_statistiques(stats)
        except Exception as erreur:
            messagebox.showerror("Erreur", f"{type(erreur).__name__} : {erreur}")
            return

        self.zone_resultat.config(text=resultat)
        # Force le reaffichage : certaines versions de Tk sur macOS ne
        # redessinent pas la fenetre en sortant d'une boite de dialogue.
        self.zone_resultat.update_idletasks()


def main():
    Application().mainloop()


if __name__ == "__main__":
    main()
